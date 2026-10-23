"""Read-only inventory of LoDIR workbooks. Writes documentation, never sources.

Run with Python and openpyxl. No workbook is recalculated or modified.
"""
import collections
import datetime
import hashlib
import json
import pathlib
import re

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parents[1]
BASE = ROOT / 'data' / 'Diagnósticos'
OUT = ROOT / 'docs' / 'diagnosticos'
OUT.mkdir(parents=True, exist_ok=True)

def serial(value):
    return value.isoformat() if isinstance(value, (datetime.datetime, datetime.date)) else value

reports = []
datasets = []
for path in sorted(BASE.rglob('*.xlsx')):
    book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sheet = book.worksheets[0]
    headers = [sheet.cell(2, c).value for c in range(1, 50)]
    records = [(i, list(row[:49])) for i, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3)
               if isinstance(row[0], str) and re.fullmatch(r'[A-Z]{1,3}-\d+', row[0].strip())]
    codes = collections.Counter(row[0].strip() for _, row in records)
    dates = collections.Counter(str(serial(row[4])) for _, row in records)
    rows = [[serial(v) for v in row] for _, row in records]
    digest = hashlib.sha256(json.dumps(rows, ensure_ascii=False).encode()).hexdigest()
    validation = book.worksheets[1]
    options = [{
        'column': get_column_letter(c), 'question': validation.cell(2, c).value,
        'options': [serial(validation.cell(r, c).value) for r in range(3, validation.max_row + 1)
                    if validation.cell(r, c).value is not None]
    } for c in range(1, validation.max_column + 1) if validation.cell(2, c).value]
    report = {
        'file': path.relative_to(ROOT).as_posix(), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'data_sha256': digest, 'sheets': [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in book],
        'records': len(records), 'unique_codes': len(codes), 'duplicate_codes': {k:v for k,v in codes.items() if v > 1},
        'first_record_row': records[0][0] if records else None, 'last_record_row': records[-1][0] if records else None,
        'dates': dict(dates), 'subprefeituras': dict(collections.Counter(row[1] for _,row in records)),
        'statuses': dict(collections.Counter(row[6] for _,row in records)),
        'missing_by_column': {get_column_letter(c+1): sum(row[c] is None for _,row in records) for c in range(49)},
        'invalid_coordinates': [i for i,row in records if not (isinstance(row[12], (int,float)) and isinstance(row[13], (int,float)) and -90 <= row[12] <= 90 and -180 <= row[13] <= 180)],
        'headers': headers, 'options': options,
        'data_validations': [{'range': str(v.sqref), 'type': v.type, 'formula': v.formula1} for v in sheet.data_validations.dataValidation],
        'formulas': sum(c.data_type == 'f' for s in openpyxl.load_workbook(path, read_only=True, data_only=False) for row in s for c in row),
    }
    reports.append(report)
    # Comparison only: retain first occurrence, while reporting every duplicate.
    indexed = {}
    for row in rows:
        indexed.setdefault(row[0], row)
    datasets.append(indexed)

features = json.loads((ROOT/'data'/'features-PV.json').read_text(encoding='utf-8'))['features']
map_codes = collections.Counter(f.get('setor','').strip() for f in features)
comparisons = []
for i in range(len(reports)):
    for j in range(i+1,len(reports)):
        shared = datasets[i].keys() & datasets[j].keys()
        comparisons.append({'a': i, 'b': j, 'shared_codes': len(shared),
                            'identical_rows': sum(datasets[i][k] == datasets[j][k] for k in shared)})
audit = {'sources': reports, 'comparisons': comparisons, 'map': {
    'features': len(features), 'unique_codes': len(map_codes),
    'duplicate_codes': {k:v for k,v in map_codes.items() if v > 1},
    'latest_missing_in_map': sorted(datasets[-1].keys() - map_codes.keys()),
    'map_missing_in_latest': sorted(map_codes.keys() - datasets[-1].keys()),
}}
(OUT/'inventario.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
lines = ['# Dicionário dos levantamentos LoDIR', '', 'Extraído das seis fontes. Colunas A:AW (49 campos). Rótulos e alternativas originais estão no inventário JSON. Não são regras novas do aplicativo.', '']
latest = reports[-1]
for c,h in enumerate(latest['headers'],1):
    lines += [f'## {get_column_letter(c)} — campo {c}', '', str(h), '']
lines += ['# Listas de respostas da versão de 19/10/2026', '']
for option in latest['options']:
    lines += [f"## Lista {option['column']}", '', str(option['question']), '']
    lines += [f'- {v}' for v in option['options']]
    lines += ['']
(OUT/'dicionario-campos.md').write_text('\n'.join(lines),encoding='utf-8')
print(json.dumps({'sources':[{k:s[k] for k in ['file','records','unique_codes','duplicate_codes','dates','statuses','formulas','invalid_coordinates']} for s in reports], 'comparisons':comparisons,'map':audit['map']}, ensure_ascii=True,indent=2))
